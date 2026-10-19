from dataclasses import dataclass
from pathlib import Path


@dataclass
class ParsedPage:
    page: int
    title: str | None
    text: str
    content_type: str = "text"


def parse_document(path: Path) -> list[ParsedPage]:
    suffix = path.suffix.lower()
    if suffix in {".txt", ".md"}:
        return _parse_text(path)
    if suffix == ".pptx":
        return _parse_pptx(path)
    if suffix == ".pdf":
        return _parse_pdf(path)
    raise ValueError(f"Unsupported file type: {suffix}")


def _parse_text(path: Path) -> list[ParsedPage]:
    try:
        text = path.read_text(encoding="utf-8")
    except UnicodeDecodeError:
        text = path.read_text(encoding="gbk")
    text = text.strip()
    if not text:
        raise ValueError("The file contains no text")
    title = next((line.strip() for line in text.splitlines() if line.strip()), None)
    return [ParsedPage(page=1, title=title, text=text)]


def _collect_shape_text(shapes, parts: list[str], group_type) -> None:
    for shape in shapes:
        if getattr(shape, "has_text_frame", False):
            text = shape.text_frame.text.strip()
            if text:
                parts.append(text)
        if getattr(shape, "has_table", False):
            rows = []
            for row in shape.table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                rows.append(" | ".join(cells))
            table_text = "\n".join(rows).strip()
            if table_text:
                parts.append(table_text)
        if shape.shape_type == group_type:
            _collect_shape_text(shape.shapes, parts, group_type)


def _parse_pptx(path: Path) -> list[ParsedPage]:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    presentation = Presentation(str(path))
    pages: list[ParsedPage] = []

    for index, slide in enumerate(presentation.slides, start=1):
        parts: list[str] = []
        _collect_shape_text(slide.shapes, parts, MSO_SHAPE_TYPE.GROUP)

        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                parts.append(f"[备注] {notes}")

        text = "\n".join(parts).strip()
        if not text:
            continue

        title = None
        title_shape = slide.shapes.title
        if title_shape is not None and getattr(title_shape, "has_text_frame", False):
            title_text = title_shape.text_frame.text.strip()
            if title_text:
                title = title_text
        pages.append(ParsedPage(page=index, title=title, text=text))

    if not pages:
        raise ValueError("No extractable text found in the PPTX file")
    return pages


def _parse_pdf(path: Path) -> list[ParsedPage]:
    import fitz

    pages: list[ParsedPage] = []
    with fitz.open(str(path)) as document:
        for index, page in enumerate(document, start=1):
            text = page.get_text("text").strip()
            if not text:
                continue
            title = next((line.strip() for line in text.splitlines() if line.strip()), None)
            pages.append(ParsedPage(page=index, title=title, text=text))

    if not pages:
        raise ValueError("No extractable text found in the PDF file")
    return pages
